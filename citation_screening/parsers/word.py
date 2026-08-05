import hashlib
import io
import re
from typing import Any, Dict, List

from docx import Document

from .context import add_context, split_candidate_sentences


_CITATION_GROUP = re.compile(r"\[(\s*\d+(?:\s*[-–—,;]\s*\d+)*)\]")


def _paragraph_text(paragraph) -> str:
    chunks: List[str] = []
    superscript_buffer = ""
    for run in paragraph.runs:
        text = run.text or ""
        is_citation_superscript = bool(
            run.font.superscript and re.fullmatch(r"[\d\s,;\-–—]+", text)
        )
        if is_citation_superscript:
            superscript_buffer += text
            continue
        if superscript_buffer:
            chunks.append(f"[{superscript_buffer.strip()}]")
            superscript_buffer = ""
        chunks.append(text)
    if superscript_buffer:
        chunks.append(f"[{superscript_buffer.strip()}]")
    return "".join(chunks).strip()


def _expand_group(group: str) -> List[str]:
    values: List[int] = []
    for part in re.split(r"\s*[,;]\s*", group.strip()):
        range_match = re.fullmatch(r"(\d+)\s*[-–—]\s*(\d+)", part)
        if range_match:
            start, end = map(int, range_match.groups())
            if 0 < start <= end and end - start <= 100:
                values.extend(range(start, end + 1))
        elif part.strip().isdigit():
            values.append(int(part.strip()))
    return [f"B{value}" for value in dict.fromkeys(values)]


def _citation_rids(text: str) -> List[str]:
    rids: List[str] = []
    for match in _CITATION_GROUP.finditer(text):
        rids.extend(_expand_group(match.group(1)))
    return list(dict.fromkeys(rids))


def _make_item(text: str, source: str, location: Dict[str, int]) -> Dict[str, Any]:
    rids = _citation_rids(text)
    return {
        "sentence_id": hashlib.md5(f"{source}:{location}:{text}".encode("utf-8")).hexdigest(),
        "sentence_text": re.sub(r"\s+", " ", text).strip(),
        "source": source,
        "location": location,
        "citations": [{"rid": rid} for rid in rids],
    }


def parse_word(file_bytes: bytes) -> List[Dict[str, Any]]:
    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise ValueError(f"无法读取 Word 文件：{exc}") from exc

    results: List[Dict[str, Any]] = []
    for paragraph_index, paragraph in enumerate(document.paragraphs, 1):
        text = _paragraph_text(paragraph)
        if re.fullmatch(r"\s*(references|bibliography|参考文献)\s*[:：]?\s*", text, re.IGNORECASE):
            break
        sentence_parts = split_candidate_sentences(text)
        for sentence_index, sentence in enumerate(sentence_parts):
            if _citation_rids(sentence):
                item = _make_item(sentence, "body_text", {"paragraph": paragraph_index})
                results.append(add_context(item, sentence_parts, sentence_index, text))

    for table_index, table in enumerate(document.tables, 1):
        for row_index, row in enumerate(table.rows, 1):
            text = " | ".join(
                filter(None, (" ".join(_paragraph_text(p) for p in cell.paragraphs).strip() for cell in row.cells))
            )
            if _citation_rids(text):
                item = _make_item(text, "table_content", {"table": table_index, "row": row_index})
                item.update({"context_before": "", "context_after": "", "full_block": text, "boundary_confidence": "table_row"})
                results.append(item)
    return results
